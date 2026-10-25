# -*- coding: utf-8 -*-
"""Build 军训心得.docx from 军训心得.md with standard Chinese homework formatting."""
import io
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = "军训心得.md"
OUT = "军训心得.docx"

TITLE_MAIN = "烈日下的第一课"
TITLE_SUB = "——军训心得"


def style_run(run, ascii_font, ea_font, size_pt, bold=False):
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), ascii_font)
    rFonts.set(qn("w:hAnsi"), ascii_font)
    rFonts.set(qn("w:eastAsia"), ea_font)
    rFonts.set(qn("w:cs"), ascii_font)


def set_first_line_chars(paragraph, chars=2, font_pt=12):
    pPr = paragraph._p.get_or_add_pPr()
    ind = OxmlElement("w:ind")
    ind.set(qn("w:firstLineChars"), str(int(chars * 100)))
    ind.set(qn("w:firstLine"), str(int(chars * font_pt * 20)))
    pPr.append(ind)


def main():
    with io.open(SRC, encoding="utf-8") as fh:
        lines = [ln.strip() for ln in fh.read().splitlines()]
    body = [ln for ln in lines if ln and not ln.startswith("#")]

    doc = Document()

    # Page setup: A4, Word "normal" Chinese margins
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(3.18)
    sec.right_margin = Cm(3.18)

    # Normal style baseline
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    # Title (line 1): 二号 黑体, centered
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    style_run(p.add_run(TITLE_MAIN), "Times New Roman", "黑体", 22, bold=False)

    # Title (line 2): 三号 楷体, centered
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.5
    style_run(p.add_run(TITLE_SUB), "Times New Roman", "楷体", 16, bold=False)

    # Body: 宋体 小四, 首行缩进 2 字符, 1.5 倍行距, 两端对齐
    for text in body:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing = 1.5
        set_first_line_chars(p, 2, 12)
        style_run(p.add_run(text), "Times New Roman", "宋体", 12)

    doc.save(OUT)
    print("saved", OUT, "paragraphs:", len(body) + 2)


if __name__ == "__main__":
    main()
